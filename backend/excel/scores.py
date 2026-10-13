import os
import json
import openpyxl
from openpyxl.comments import Comment
from openpyxl.styles import PatternFill, Font
from pathlib import Path
from typing import Optional

from excel.locking import locked

_default_data = Path(__file__).parent.parent.parent / "data"
DATA_DIR = Path(os.environ.get("DATA_DIR", str(_default_data)))
SCORES_PATH = DATA_DIR / "scores.xlsx"
DARK_PURPLE = "FF674EA7"
LIGHT_PURPLE = "FFB4A7D6"


_NON_PERSON_COLS = {
    "ile osob", "ile osob per spotkanie", "średnia filmów", "srednia filmow",
    "sredni hejt", "srednia oceniania", "s", "ile osob ", "średnia filmów ",
    "l", "srednia", "ilość obejrzanych filmów", "ilosc obejrzanych filmow",
}


def _is_formula(v) -> bool:
    return isinstance(v, str) and v.startswith("=")


def _is_person_col(name: str) -> bool:
    return name.lower().strip() not in _NON_PERSON_COLS


def _row_has_scores(row, people_cols: dict) -> bool:
    return any(
        isinstance(row[col - 1].value, (int, float))
        for col in people_cols
        if col - 1 < len(row)
    )


def _row_metadata(row) -> dict:
    comment = row[0].comment
    if comment and comment.author == "Movie Night":
        try:
            value = json.loads(comment.text)
            return value if isinstance(value, dict) else {}
        except ValueError:
            pass
    return {}


def _is_movie_row(row, people_cols: dict) -> bool:
    # New films remain films even if everyone skipped scoring them.
    return _row_metadata(row).get("type") == "movie" or _row_has_scores(row, people_cols)


def _extract_scores(row, people_cols: dict) -> dict:
    scores = {}
    for col, name in people_cols.items():
        if col - 1 < len(row):
            v = row[col - 1].value
            if isinstance(v, (int, float)):
                scores[name] = float(v)
    return scores


def _get_avg(row, scores: dict) -> Optional[float]:
    v = row[1].value
    if isinstance(v, (int, float)):
        return round(float(v), 2)
    vals = list(scores.values())
    return round(sum(vals) / len(vals), 2) if vals else None


def _get_people_cols(ws, header_row: int) -> dict:
    people = {}
    for cell in ws[header_row][2:]:
        v = cell.value
        if v and isinstance(v, str) and not _is_formula(v) and _is_person_col(v):
            people[cell.column] = v.strip()
    return people


def _find_header_row(ws, candidates=("xero", "kacper", "mati", "wiktoria")) -> int:
    for r in ws.iter_rows(min_row=1, max_row=6):
        for cell in r:
            if cell.value and str(cell.value).strip().lower() in candidates:
                return cell.row
    return 1


def parse_season4(wb) -> list:
    ws = wb["Sezon 4 - Oceny"]
    people_cols = _get_people_cols(ws, 1)
    current_category = None
    movies = []
    for row in ws.iter_rows(min_row=4):
        title = row[0].value
        if not title or _is_formula(str(title)):
            continue
        title = str(title).strip()
        if not title:
            continue
        if not _is_movie_row(row, people_cols):
            current_category = title
            continue
        scores = _extract_scores(row, people_cols)
        movies.append({
            "title": title,
            "category": _row_metadata(row).get("category") or current_category,
            "season": 4,
            "average": _get_avg(row, scores),
            "scores": scores,
        })
    return movies


def parse_season3(wb) -> list:
    ws = wb["Sezon 3 - Oceny"]
    header_row = _find_header_row(ws)
    people_cols = _get_people_cols(ws, header_row)
    current_category = None
    movies = []
    for row in ws.iter_rows(min_row=header_row + 1):
        title = row[0].value
        if not title or _is_formula(str(title)):
            continue
        title = str(title).strip()
        if not title:
            continue
        if title.lower() in ("średnia nocy", "srednia nocy"):
            continue
        if not _is_movie_row(row, people_cols):
            current_category = title
            continue
        scores = _extract_scores(row, people_cols)
        movies.append({
            "title": title,
            "category": current_category,
            "season": 3,
            "average": _get_avg(row, scores),
            "scores": scores,
        })
    return movies


def parse_season2(wb) -> list:
    ws = wb["Sezon 2 - Oceny"]
    people_cols = _get_people_cols(ws, 1)
    movies = []
    for row in ws.iter_rows(min_row=6):
        title = row[0].value
        if not title or _is_formula(str(title)):
            continue
        title = str(title).strip()
        if not title or title.lower() in ("średnia nocy", "srednia nocy"):
            continue
        if not _is_movie_row(row, people_cols):
            continue
        scores = _extract_scores(row, people_cols)
        movies.append({
            "title": title,
            "category": None,
            "season": 2,
            "average": _get_avg(row, scores),
            "scores": scores,
        })
    return movies


def parse_season1(wb) -> list:
    ws = wb["Sezon 1 - Oceny"]
    people_cols = _get_people_cols(ws, 1)
    movies = []
    for row in ws.iter_rows(min_row=3):
        title = row[0].value
        if not title or _is_formula(str(title)):
            continue
        title = str(title).strip()
        if not title:
            continue
        if not _is_movie_row(row, people_cols):
            continue
        scores = _extract_scores(row, people_cols)
        movies.append({
            "title": title,
            "category": None,
            "season": 1,
            "average": _get_avg(row, scores),
            "scores": scores,
        })
    return movies


@locked(lambda: SCORES_PATH)
def get_all_header_people() -> list[str]:
    """Return all raw person names from all seasons' column headers."""
    wb = openpyxl.load_workbook(SCORES_PATH, data_only=True)
    seen: dict[str, str] = {}
    for sheet_name in _SHEET_NAMES.values():
        ws = wb[sheet_name]
        header_row = _find_header_row(ws) if sheet_name == "Sezon 3 - Oceny" else 1
        for raw in _get_people_cols(ws, header_row).values():
            key = raw.strip().lower()
            if key not in seen:
                seen[key] = raw.strip()
    return list(seen.values())


@locked(lambda: SCORES_PATH)
def get_all_movies() -> list:
    wb = openpyxl.load_workbook(SCORES_PATH, data_only=True)
    result = []
    result.extend(parse_season1(wb))
    result.extend(parse_season2(wb))
    result.extend(parse_season3(wb))
    result.extend(parse_season4(wb))
    return result


def _parse_nights_from_sheet(ws, people_cols: dict, start_row: int, has_categories: bool) -> list:
    nights = []
    current_name = None
    current_movies = []
    current_hosts = None

    for row in ws.iter_rows(min_row=start_row):
        title = row[0].value
        if title is None or (isinstance(title, str) and not title.strip()):
            # Blank row: split night for non-category seasons
            if not has_categories and current_movies:
                nights.append({"name": current_name, "movies": current_movies, "hosts": current_hosts})
                current_movies = []
                current_name = None
                current_hosts = None
            continue
        if _is_formula(str(title)):
            continue
        title = str(title).strip()

        if title.lower() in ("średnia nocy", "srednia nocy"):
            if current_movies:
                nights.append({"name": current_name, "movies": current_movies, "hosts": current_hosts})
            current_movies = []
            current_name = None
            current_hosts = None
            continue

        if not _is_movie_row(row, people_cols):
            if has_categories:
                if current_movies:
                    nights.append({"name": current_name, "movies": current_movies, "hosts": current_hosts})
                current_name = title
                current_hosts = _row_metadata(row).get("hosts")
                current_movies = []
            continue

        scores = _extract_scores(row, people_cols)
        avg = _get_avg(row, scores)
        current_movies.append({"title": title, "category": _row_metadata(row).get("category") or current_name, "average": avg, "scores": scores, "row_num": row[0].row})

    if current_movies:
        nights.append({"name": current_name, "movies": current_movies, "hosts": current_hosts})

    return nights


@locked(lambda: SCORES_PATH)
def get_all_nights() -> list:
    wb = openpyxl.load_workbook(SCORES_PATH, data_only=True)
    result = []

    # Season 1
    ws1 = wb["Sezon 1 - Oceny"]
    people_cols1 = _get_people_cols(ws1, 1)
    for night in _parse_nights_from_sheet(ws1, people_cols1, start_row=3, has_categories=False):
        night["season"] = 1
        result.append(night)

    # Season 2
    ws2 = wb["Sezon 2 - Oceny"]
    people_cols2 = _get_people_cols(ws2, 1)
    for night in _parse_nights_from_sheet(ws2, people_cols2, start_row=6, has_categories=False):
        night["season"] = 2
        result.append(night)

    # Season 3
    ws3 = wb["Sezon 3 - Oceny"]
    header_row3 = _find_header_row(ws3)
    people_cols3 = _get_people_cols(ws3, header_row3)
    for night in _parse_nights_from_sheet(ws3, people_cols3, start_row=header_row3 + 1, has_categories=True):
        night["season"] = 3
        result.append(night)

    # Season 4
    ws4 = wb["Sezon 4 - Oceny"]
    people_cols4 = _get_people_cols(ws4, 1)
    for night in _parse_nights_from_sheet(ws4, people_cols4, start_row=4, has_categories=True):
        night["season"] = 4
        result.append(night)

    return result


@locked(lambda: SCORES_PATH)
def get_last_night_people() -> list:
    wb = openpyxl.load_workbook(SCORES_PATH, data_only=True)
    ws4 = wb["Sezon 4 - Oceny"]
    people_cols4 = _get_people_cols(ws4, 1)
    nights = _parse_nights_from_sheet(ws4, people_cols4, start_row=4, has_categories=True)
    if not nights:
        return []
    last = nights[-1]
    scorers = set()
    for movie in last["movies"]:
        for person in movie.get("scores", {}).keys():
            scorers.add(person)
    return sorted(scorers)


@locked(lambda: SCORES_PATH)
def get_season4_people() -> list:
    wb = openpyxl.load_workbook(SCORES_PATH, data_only=True)
    ws = wb["Sezon 4 - Oceny"]
    seen: dict[str, str] = {}
    for cell in ws[1][2:]:
        v = cell.value
        if v and isinstance(v, str) and not _is_formula(v) and _is_person_col(v):
            key = v.strip().lower()
            if key not in seen:
                seen[key] = v.strip()
    return list(seen.values())


_SHEET_NAMES = {
    1: "Sezon 1 - Oceny",
    2: "Sezon 2 - Oceny",
    3: "Sezon 3 - Oceny",
    4: "Sezon 4 - Oceny",
}


@locked(lambda: SCORES_PATH)
def get_people_for_season(season: int) -> list[str]:
    wb = openpyxl.load_workbook(SCORES_PATH, data_only=True)
    ws = wb[_SHEET_NAMES[season]]
    header_row = _find_header_row(ws) if season == 3 else 1
    return list(_get_people_cols(ws, header_row).values())


def _require_movie_row(ws, season: int, row_num: int, people_cols: dict):
    start_row = {1: 3, 2: 6, 3: _find_header_row(ws) + 1, 4: 4}[season]
    if row_num < start_row or row_num > ws.max_row:
        raise ValueError("Movie row not found")
    row = ws[row_num]
    if not row[0].value or _is_formula(row[0].value) or not _is_movie_row(row, people_cols):
        raise ValueError("Movie row not found")
    return row


@locked(lambda: SCORES_PATH)
def edit_score(season: int, row_num: int, person_canonical: str, score: Optional[float]):
    from config import normalize_name  # local import avoids circular at module level
    wb = openpyxl.load_workbook(SCORES_PATH)
    ws = wb[_SHEET_NAMES[season]]
    header_row = _find_header_row(ws) if season == 3 else 1
    people_cols = _get_people_cols(ws, header_row)
    row = _require_movie_row(ws, season, row_num, people_cols)

    col = next((c for c, raw in people_cols.items() if normalize_name(raw) == person_canonical), None)
    if col is None:
        raise ValueError(f"Person '{person_canonical}' not found in season {season} sheet")

    # Preserve a legacy film even when its final remaining score is removed.
    metadata = {**_row_metadata(row), "type": "movie"}
    row[0].comment = Comment(json.dumps(metadata, ensure_ascii=False), "Movie Night")

    light_fill = PatternFill(start_color=LIGHT_PURPLE, end_color=LIGHT_PURPLE, fill_type="solid")
    cell = ws.cell(row_num, col)
    cell.value = score
    if score is not None:
        cell.fill = light_fill

    row_cells = ws[row_num]
    scores_in_row = _extract_scores(row_cells, people_cols)
    new_avg = round(sum(scores_in_row.values()) / len(scores_in_row), 2) if scores_in_row else None
    ws.cell(row_num, 2).value = new_avg

    wb.save(SCORES_PATH)


@locked(lambda: SCORES_PATH)
def delete_night(season: int, night_name: str):
    wb = openpyxl.load_workbook(SCORES_PATH)
    ws = wb[_SHEET_NAMES[season]]
    header_row_num = _find_header_row(ws) if season == 3 else 1
    people_cols = _get_people_cols(ws, header_row_num)
    start_row = {1: 3, 2: 6, 3: header_row_num + 1, 4: 4}[season]
    has_categories = season in (3, 4)

    rows_to_delete: list[int] = []
    in_target = False

    for row in ws.iter_rows(min_row=start_row):
        raw = row[0].value
        if raw is None or (isinstance(raw, str) and not raw.strip()):
            if not has_categories and in_target:
                break
            continue
        if _is_formula(str(raw)):
            continue
        title = str(raw).strip()

        if has_categories:
            if not _is_movie_row(row, people_cols):
                if in_target:
                    break  # next night starts, stop
                if title == night_name:
                    in_target = True
                    rows_to_delete.append(row[0].row)
            elif in_target:
                rows_to_delete.append(row[0].row)

    if not rows_to_delete:
        raise ValueError(f"Night '{night_name}' not found in season {season}")

    for r in sorted(rows_to_delete, reverse=True):
        ws.delete_rows(r)

    wb.save(SCORES_PATH)


@locked(lambda: SCORES_PATH)
def delete_movie_row(season: int, row_num: int):
    wb = openpyxl.load_workbook(SCORES_PATH)
    ws = wb[_SHEET_NAMES[season]]
    header_row = _find_header_row(ws) if season == 3 else 1
    _require_movie_row(ws, season, row_num, _get_people_cols(ws, header_row))
    ws.delete_rows(row_num)
    wb.save(SCORES_PATH)


@locked(lambda: SCORES_PATH)
def add_person_to_scores(name: str) -> bool:
    """Add a season-4 column for a person.

    Returns True if a column was created, False if one already existed. Without
    this guard a repeated "Dodaj" creates a second column for the same person,
    splitting their scores across both.
    """
    wb = openpyxl.load_workbook(SCORES_PATH)
    ws = wb["Sezon 4 - Oceny"]

    from config import normalize_name
    target = normalize_name(name)
    last_person_col = 2
    for cell in ws[1]:
        if cell.column >= 3 and cell.value and isinstance(cell.value, str):
            if not _is_formula(cell.value) and normalize_name(cell.value) == target:
                return False
            last_person_col = cell.column

    new_col = last_person_col + 1
    ws.cell(1, new_col, value=name)
    dark_fill = PatternFill(start_color=DARK_PURPLE, end_color=DARK_PURPLE, fill_type="solid")
    ws.cell(1, new_col).fill = dark_fill
    wb.save(SCORES_PATH)
    return True


@locked(lambda: SCORES_PATH)
def _session_columns(ws) -> tuple[dict, int]:
    """{person name -> column} for the season-4 header, plus the last column.

    Both the raw header ("bart") and its canonical form ("Bartosz") are keys,
    because /scores/people now hands out canonical names while the sheet still
    stores whatever was typed into the header years ago.
    """
    from config import normalize_name  # local import avoids circular at module level

    people_cols: dict[str, int] = {}
    max_col = 2
    for cell in ws[1][2:]:
        if not (cell.value and isinstance(cell.value, str)) or _is_formula(cell.value):
            continue
        raw = cell.value.strip()
        max_col = max(max_col, cell.column)
        if not _is_person_col(raw):
            continue
        people_cols.setdefault(raw, cell.column)
        people_cols.setdefault(normalize_name(raw), cell.column)
    return people_cols, max_col


@locked(lambda: SCORES_PATH)
def save_session(
    movies: list,
    category: Optional[str],
    scores: dict,
    present_people: list,
    movie_categories: dict | None = None,
    hosts: list[str] | None = None,
):
    wb = openpyxl.load_workbook(SCORES_PATH)
    ws = wb["Sezon 4 - Oceny"]

    people_cols, max_person_col = _session_columns(ws)
    if max_person_col <= 2:
        max_person_col = 20

    # Only people marked present get a score written. Without this a stale
    # score left in the payload for someone who went home would still land in
    # the sheet.
    from config import normalize_name
    present = {normalize_name(p) for p in present_people}

    # Find last used row
    last_row = ws.max_row
    while last_row > 3 and all(ws.cell(last_row, c).value is None for c in range(1, max_person_col + 1)):
        last_row -= 1
    next_row = last_row + 1

    dark_fill = PatternFill(start_color=DARK_PURPLE, end_color=DARK_PURPLE, fill_type="solid")
    light_fill = PatternFill(start_color=LIGHT_PURPLE, end_color=LIGHT_PURPLE, fill_type="solid")
    white_font = Font(color="FFFFFFFF")

    # Every submission is a night, including one without a theme. Otherwise
    # its films and attendance would be merged into the preceding night.
    for c in range(1, max_person_col + 1):
        cell = ws.cell(next_row, c)
        cell.fill = dark_fill
        cell.font = white_font
    ws.cell(next_row, 1).value = category or "Bez tematu"
    ws.cell(next_row, 1).comment = Comment(json.dumps({"type": "night", "hosts": hosts or []}, ensure_ascii=False), "Movie Night")
    next_row += 1

    for movie_title in movies:
        movie_scores = scores.get(movie_title, {})
        movie_scores = {normalize_name(p): v for p, v in movie_scores.items() if normalize_name(p) in present}
        movie_scores = {p: v for p, v in movie_scores.items() if p in people_cols}
        # Aliases can address the same column; average exactly what is written.
        column_scores = {people_cols[p]: v for p, v in movie_scores.items() if v is not None}
        valid = list(column_scores.values())
        avg = round(sum(valid) / len(valid), 2) if valid else None

        for c in range(1, max_person_col + 1):
            ws.cell(next_row, c).fill = light_fill

        ws.cell(next_row, 1).value = movie_title
        ws.cell(next_row, 1).comment = Comment(json.dumps({
            "type": "movie", "category": (movie_categories or {}).get(movie_title, category),
        }, ensure_ascii=False), "Movie Night")
        if avg is not None:
            ws.cell(next_row, 2).value = avg

        for col, score in column_scores.items():
            ws.cell(next_row, col).value = score

        next_row += 1

    wb.save(SCORES_PATH)

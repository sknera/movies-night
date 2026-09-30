"""Builds a miniature stand-in for data/ so tests never touch the real books.

The layout mirrors the quirks the parsers actually cope with in production:
a raw header ("bart") whose canonical name is different, a hidden person, an
aggregate column that is not a person, and one season whose header row is not
row 1.
"""
import json
from pathlib import Path

import bcrypt
import openpyxl

ADMIN_PASSWORD = "test-password"

# Season-4 people, as they appear in the sheet header.
S4_HEADERS = ["Xero", "Kaja", "bart", "Stary"]

WHEEL_COLUMNS = {
    "xero": ["Horrory", "Kosmos", "Dokumenty"],
    "kaja": ["Anime", "Musicale"],
    "bart": ["Western"],
    "stary": ["Stare kino"],
}

# One past spin, for a category nobody still has listed.
WHEEL_HISTORY = [("kaja", "Retro")]


def _write_season4(wb):
    ws = wb.create_sheet("Sezon 4 - Oceny") if "Sezon 4 - Oceny" not in wb.sheetnames else wb["Sezon 4 - Oceny"]
    ws.cell(1, 2, "Średnia")
    for i, name in enumerate(S4_HEADERS):
        ws.cell(1, 3 + i, name)
    ws.cell(1, 3 + len(S4_HEADERS), "Ile osob")  # aggregate, not a person

    # Row 4 onwards: a category row followed by its movies.
    ws.cell(4, 1, "Kategoria startowa")
    ws.cell(5, 1, "Stary Film")
    ws.cell(5, 2, 6.0)
    for i, score in enumerate([6.0, 7.0, 5.0, 6.0]):
        ws.cell(5, 3 + i, score)
    return ws


def _write_plain_season(wb, title, header_row, start_row, movie_rows):
    ws = wb.create_sheet(title)
    ws.cell(header_row, 2, "Średnia")
    for i, name in enumerate(S4_HEADERS[:3]):
        ws.cell(header_row, 3 + i, name)
    for offset, (movie, scores) in enumerate(movie_rows):
        r = start_row + offset
        ws.cell(r, 1, movie)
        ws.cell(r, 2, round(sum(scores) / len(scores), 2))
        for i, score in enumerate(scores):
            ws.cell(r, 3 + i, score)
    return ws


def build_scores(path: Path):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    _write_season4(wb)
    # Season 3's header row is found by looking for a known name, not by position.
    _write_plain_season(wb, "Sezon 3 - Oceny", header_row=2, start_row=3,
                        movie_rows=[("Trzeci Film", [7.0, 8.0, 6.0])])
    _write_plain_season(wb, "Sezon 2 - Oceny", header_row=1, start_row=6,
                        movie_rows=[("Drugi Film", [5.0, 5.5, 6.0])])
    _write_plain_season(wb, "Sezon 1 - Oceny", header_row=1, start_row=3,
                        movie_rows=[("Pierwszy Film", [8.0, 7.5, 9.0])])
    wb.save(path)


def build_spread(path: Path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheetgo_losowanie"

    cols = list(WHEEL_COLUMNS)
    for i, key in enumerate(cols):
        ws.cell(1, 1 + i, key)
        for r, cat in enumerate(WHEEL_COLUMNS[key], start=2):
            ws.cell(r, 1 + i, cat)

    dupa_col = len(cols) + 1
    trafione_col = dupa_col + 1
    ws.cell(1, dupa_col, "dupa")
    ws.cell(1, trafione_col, "trafione")
    ws.cell(1, trafione_col + 1, "kategoria")
    for r, (person, cat) in enumerate(WHEEL_HISTORY, start=2):
        ws.cell(r, trafione_col, person)
        ws.cell(r, trafione_col + 1, cat)

    wb.save(path)


def build_config(path: Path):
    config = {
        "admin_hash": bcrypt.hashpw(ADMIN_PASSWORD.encode(), bcrypt.gensalt()).decode(),
        "hidden_people": ["Stary"],
        "name_map": {
            "xero": "Xero",
            "kaja": "Kaja",
            "bart": "Bartosz",
            "bartosz": "Bartosz",
            "stary": "Stary",
        },
        "person_colors": {
            "Xero": "#d84b4b",
            "Kaja": "#4bd892",
            "Bartosz": "#6e4bd8",
        },
        "category_hosts": {},
        "next_night_date": "2030-01-01",
    }
    path.write_text(json.dumps(config, indent=2, ensure_ascii=False), encoding="utf-8")


def build_all(data_dir: Path):
    data_dir.mkdir(parents=True, exist_ok=True)
    for stale in data_dir.glob("*.lock"):
        stale.unlink()
    build_scores(data_dir / "scores.xlsx")
    build_spread(data_dir / "spread.xlsx")
    build_config(data_dir / "config.json")
